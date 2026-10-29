"""Core analysis for N2O frontal chemisorption (reactive frontal chromatography) on reduced Cu catalysts.

Reaction:  2 Cu_s + N2O -> Cu2O + N2
The N2O that goes missing at the reactor outlet while the Cu surface reacts gives the number of surface Cu atoms.

Works on PV MassSpec "Plot Data Export / Bin Trend" .dat files (one block per curve, columns Curve / X / Y),
and on plain CSV / Excel files with a time column and one column per signal.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field, asdict

import numpy as np
import pandas as pd

N_AVOGADRO = 6.02214076e23
M_CU, M_O = 63.546, 15.999

LOADING_BASES = (
    "wt% Cu in reduced catalyst",
    "wt% CuO in calcined catalyst",
    "wt% Cu in calcined catalyst",
)


# ----------------------------------------------------------------------------- reading
def read_pv_massspec_dat(data: bytes | str) -> pd.DataFrame:
    """Parse a PV MassSpec .dat export. Returns a DataFrame: time + one column per curve ('curve_1', ...)."""
    text = data.decode("utf-8", errors="replace") if isinstance(data, bytes) else data
    curves: dict[int, list[tuple[float, float]]] = {}
    for line in text.splitlines():
        parts = line.split("\t")
        if len(parts) < 3 or not parts[0].strip().isdigit() or not parts[1].strip():
            continue
        try:
            curves.setdefault(int(parts[0]), []).append((float(parts[1]), float(parts[2])))
        except ValueError:
            continue
    if not curves:
        raise ValueError("No curve data found – is this a PV MassSpec 'Plot Data Export' file?")
    frames = []
    for k in sorted(curves):
        s = pd.DataFrame(curves[k], columns=["time", f"curve_{k}"]).drop_duplicates("time").set_index("time")
        frames.append(s)
    df = pd.concat(frames, axis=1, join="inner").reset_index()
    return df.sort_values("time").reset_index(drop=True)


def read_table(data: bytes, filename: str) -> pd.DataFrame:
    """CSV / TXT / XLSX with a header row. First column = time (s)."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xls")):
        df = pd.read_excel(io.BytesIO(data))
    else:
        df = pd.read_csv(io.BytesIO(data), sep=None, engine="python")
    df = df.apply(pd.to_numeric, errors="coerce").dropna(how="all", axis=1).dropna()
    df = df.rename(columns={df.columns[0]: "time"})
    return df.sort_values("time").reset_index(drop=True)


def load_file(data: bytes, filename: str) -> pd.DataFrame:
    if filename.lower().endswith(".dat"):
        return read_pv_massspec_dat(data)
    try:
        return read_pv_massspec_dat(data)
    except ValueError:
        return read_table(data, filename)


def guess_columns(df: pd.DataFrame) -> dict[str, str | None]:
    """Guess which column is N2, N2O and the inert reference.

    N2 has a sharp peak right after the switch; N2O rises late to a plateau; the reference is the
    largest and most constant signal. Defaults follow the usual export order (1 = N2, 2 = ref, 3 = N2O).
    """
    cols = [c for c in df.columns if c != "time"]
    if len(cols) >= 3 and {"curve_1", "curve_2", "curve_3"} <= set(cols):
        return {"n2": "curve_1", "ref": "curve_2", "n2o": "curve_3"}
    if len(cols) >= 3:
        means = df[cols].mean()
        ref = means.idxmax()
        rest = [c for c in cols if c != ref]
        # N2 peaks early: compare time of maximum
        tmax = {c: df["time"][df[c].idxmax()] for c in rest}
        n2 = min(tmax, key=tmax.get)
        n2o = [c for c in rest if c != n2][0]
        return {"n2": n2, "ref": ref, "n2o": n2o}
    if len(cols) == 2:
        return {"n2": cols[0], "ref": None, "n2o": cols[1]}
    return {"n2": None, "ref": None, "n2o": cols[0] if cols else None}


# ----------------------------------------------------------------------------- settings
@dataclass
class Sample:
    name: str = "sample"
    mass_ox_g: float = 0.3086
    cu_loading: float = 0.80
    loading_basis: str = LOADING_BASES[0]
    y_n2o: float = 0.02
    flow_ml_min: float = 50.0
    molar_volume_ml: float = 22400.0
    site_density_m2: float = 1.47e19
    stoich: float = 2.0
    rho_cu: float = 8.92
    t_red_c: float | None = None


@dataclass
class Windows:
    baseline: tuple[float, float] | None = None   # before N2O switch
    start: float | None = None                    # integration start (N2 onset)
    plateau: tuple[float, float] | None = None    # N2O plateau; plateau start = integration end
    onset_threshold: float = 0.03                 # fraction of N2 peak height
    plateau_fraction: float = 0.97                # plateau reached when smoothed N2O >= this x final level
    plateau_length_s: float = 240.0


@dataclass
class Result:
    windows: Windows
    t_equiv_n2o_s: float
    t_equiv_n2_s: float | None
    n_n2o_mol: float
    n_cus_mol: float
    s_total_m2: float
    m_cu_g: float
    m_red_g: float
    s_per_g_ox: float
    s_per_g_red: float
    s_per_g_cu: float
    dispersion: float
    d_particle_nm: float
    fragment_ratio: float | None
    plateau_noise: float
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> dict:
        d = asdict(self)
        w = d.pop("windows")
        d.pop("warnings")
        d.update({
            "integration_start_s": w["start"],
            "integration_end_s": w["plateau"][0],
            "plateau_end_s": w["plateau"][1],
            "baseline_start_s": w["baseline"][0],
            "baseline_end_s": w["baseline"][1],
        })
        return d


# ----------------------------------------------------------------------------- maths
def _mean(t, y, lo, hi):
    m = (t >= lo) & (t <= hi)
    if not m.any():
        raise ValueError(f"No data points between {lo} and {hi} s")
    return float(np.mean(y[m]))


def _trapz(t, y, lo, hi):
    m = (t >= lo) & (t <= hi)
    if m.sum() < 2:
        return 0.0
    tt, yy = t[m], y[m]
    return float(np.sum((yy[1:] + yy[:-1]) / 2 * np.diff(tt)))


def cu_masses(s: Sample) -> tuple[float, float]:
    """Return (Cu mass, reduced sample mass) in g."""
    x, m = s.cu_loading, s.mass_ox_g
    r = (M_CU + M_O) / M_CU
    if s.loading_basis == "wt% CuO in calcined catalyst":
        m_cu = x * m / r
    elif s.loading_basis == "wt% Cu in calcined catalyst":
        m_cu = x * m
    else:  # wt% Cu in reduced catalyst
        m_zr = m / (1 + x / (1 - x) * r)
        m_cu = x / (1 - x) * m_zr
    m_red = m - m_cu * M_O / M_CU
    return m_cu, m_red


def prepare(df: pd.DataFrame, cols: dict, normalize: bool = True) -> pd.DataFrame:
    out = pd.DataFrame({"time": df["time"].to_numpy(float)})
    ref = df[cols["ref"]].to_numpy(float) if (normalize and cols.get("ref")) else 1.0
    out["n2o"] = df[cols["n2o"]].to_numpy(float) / ref
    out["n2"] = df[cols["n2"]].to_numpy(float) / ref if cols.get("n2") else np.nan
    return out


def auto_windows(p: pd.DataFrame, w: Windows | None = None) -> Windows:
    """Fill in any window left as None."""
    w = w or Windows()
    t, n2, n2o = p["time"].to_numpy(), p["n2"].to_numpy(), p["n2o"].to_numpy()
    have_n2 = np.isfinite(n2).all()
    sig = n2 if have_n2 else n2o

    # rough onset from first 5 % of points
    head = max(10, len(t) // 20)
    base0, sd0 = np.mean(sig[:head]), np.std(sig[:head])
    rough = t[np.argmax(sig > base0 + max(10 * sd0, 0.05 * (sig.max() - base0)))]

    baseline = w.baseline or (float(t[0]), float(rough - 10))
    b = _mean(t, sig, *baseline)
    start = w.start
    if start is None:
        thr = b + w.onset_threshold * (sig[t > baseline[1]].max() - b)
        idx = np.where((sig > thr) & (t > baseline[1]))[0]
        start = float(t[idx[0]]) if len(idx) else float(rough)

    plateau = w.plateau
    if plateau is None:
        after = t > start
        smooth = pd.Series(n2o).rolling(15, center=True, min_periods=1).mean().to_numpy()
        tail = after & (t >= t[after][int(0.6 * after.sum())])           # last 40 % after onset
        level = np.median(n2o[tail])
        b44 = _mean(t, n2o, *baseline)
        reach = np.where(after & (smooth - b44 >= w.plateau_fraction * (level - b44)))[0]
        p0 = float(t[reach[0]]) if len(reach) else float(t[after][len(t[after]) // 2])
        p1 = min(p0 + w.plateau_length_s, float(t[-1]))
        plateau = (p0, p1)
    return Windows(baseline, start, plateau, w.onset_threshold, w.plateau_fraction, w.plateau_length_s)


def analyse(p: pd.DataFrame, s: Sample, w: Windows) -> tuple[Result, pd.DataFrame]:
    t = p["time"].to_numpy()
    n2o, n2 = p["n2o"].to_numpy(), p["n2"].to_numpy()
    warn = []
    b44 = _mean(t, n2o, *w.baseline)
    p44 = _mean(t, n2o, *w.plateau)
    if p44 <= b44:
        raise ValueError("N2O plateau is not above the baseline – check the windows / column choice.")
    f_n2o = (n2o - b44) / (p44 - b44)
    deficit = 1 - f_n2o
    t_eq = _trapz(t, deficit, w.start, w.plateau[0])
    pl = (t >= w.plateau[0]) & (t <= w.plateau[1])
    noise = float(np.std(f_n2o[pl]))

    calc = pd.DataFrame({"time": t, "n2o_norm": n2o, "f_n2o": f_n2o, "deficit": deficit})
    t_eq_n2, k = None, None
    if np.isfinite(n2).all():
        b28 = _mean(t, n2, *w.baseline)
        k = (_mean(t, n2, *w.plateau) - b28) / (p44 - b44)
        net = n2 - b28 - k * (n2o - b44)
        sel = (t >= w.start) & (t <= w.plateau[0]) & (f_n2o >= 0.05) & (f_n2o <= 0.9)
        if sel.sum() >= 5:
            c_n2 = float(np.mean(net[sel] / (1 - f_n2o[sel])))
            f_n2 = net / c_n2
            t_eq_n2 = _trapz(t, f_n2, w.start, w.plateau[0])
            calc["f_n2"] = f_n2
            calc["mass_balance"] = f_n2o + f_n2
            if abs(t_eq_n2 / t_eq - 1) > 0.15:
                warn.append(f"N2 and N2O areas differ by {100*(t_eq_n2/t_eq-1):+.0f} % – check windows.")
        else:
            warn.append("Too few points in the breakthrough to cross-check with N2.")

    if noise > 0.1:
        warn.append(f"N2O plateau is noisy ({100*noise:.0f} %) – choose a flatter plateau window.")
    if w.start <= w.baseline[1]:
        warn.append("Integration starts inside the baseline window.")

    n_flow = s.y_n2o * s.flow_ml_min / 60 / s.molar_volume_ml
    n_n2o = t_eq * n_flow
    n_cus = s.stoich * n_n2o
    s_tot = n_cus * N_AVOGADRO / s.site_density_m2
    m_cu, m_red = cu_masses(s)
    if m_cu * (M_CU + M_O) / M_CU > s.mass_ox_g * 1.0001:
        warn.append("Cu loading/basis implies more CuO than the sample mass – check the loading basis.")
    s_cu = s_tot / m_cu
    res = Result(
        windows=w, t_equiv_n2o_s=t_eq, t_equiv_n2_s=t_eq_n2, n_n2o_mol=n_n2o, n_cus_mol=n_cus,
        s_total_m2=s_tot, m_cu_g=m_cu, m_red_g=m_red, s_per_g_ox=s_tot / s.mass_ox_g,
        s_per_g_red=s_tot / m_red, s_per_g_cu=s_cu, dispersion=n_cus / (m_cu / M_CU),
        d_particle_nm=6000 / (s.rho_cu * s_cu), fragment_ratio=k, plateau_noise=noise, warnings=warn,
    )
    return res, calc


# ----------------------------------------------------------------------------- Excel report
def excel_report(raw: pd.DataFrame, cols: dict, calc: pd.DataFrame, s: Sample, r: Result) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import Reference, ScatterChart, Series
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    navy = "1F3864"
    hf, kf = PatternFill("solid", fgColor=navy), PatternFill("solid", fgColor="C6E0B4")
    box = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 12
    ws["A1"] = f"N2O chemisorption – {s.name}"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color=navy)
    rows = [
        ("Sample", None, None, "h"),
        ("Mass loaded (oxidized)", s.mass_ox_g, "g", ""),
        ("Cu loading", s.cu_loading, s.loading_basis, "%"),
        ("N2O in feed", s.y_n2o, "fraction", "%"),
        ("Flow", s.flow_ml_min, "mL/min", ""),
        ("Molar volume", s.molar_volume_ml, "mL/mol", ""),
        ("Cu site density", s.site_density_m2, "atoms/m²", "e"),
        ("Windows", None, None, "h"),
        ("Baseline", f"{r.windows.baseline[0]:.1f} – {r.windows.baseline[1]:.1f}", "s", ""),
        ("Integration", f"{r.windows.start:.1f} – {r.windows.plateau[0]:.1f}", "s", ""),
        ("N2O plateau", f"{r.windows.plateau[0]:.1f} – {r.windows.plateau[1]:.1f}", "s", ""),
        ("Results", None, None, "h"),
        ("Equivalent time of full N2O uptake", r.t_equiv_n2o_s, "s", "2"),
        ("Cross-check from N2 produced", r.t_equiv_n2_s, "s", "2"),
        ("N2O consumed", r.n_n2o_mol, "mol", "e"),
        ("Surface Cu", r.n_cus_mol, "mol", "e"),
        ("S Cu per g catalyst (oxidized)", r.s_per_g_ox, "m²/g", "k"),
        ("S Cu per g catalyst (reduced)", r.s_per_g_red, "m²/g", "k"),
        ("S Cu per g Cu", r.s_per_g_cu, "m²/g Cu", "k"),
        ("Dispersion", r.dispersion, "", "k%"),
        ("Particle size (spheres)", r.d_particle_nm, "nm", "k"),
    ]
    row = 3
    for lab, val, unit, kind in rows:
        a = ws.cell(row, 1, lab)
        a.font = Font(name="Arial", bold=kind == "h", color="FFFFFF" if kind == "h" else "000000")
        if kind == "h":
            for c in range(1, 4):
                ws.cell(row, c).fill = hf
        else:
            b = ws.cell(row, 2, val)
            ws.cell(row, 3, unit).font = Font(name="Arial", color="595959")
            b.font = Font(name="Arial", bold="k" in kind)
            b.border = box
            b.number_format = {"%": "0.00%", "k%": "0.00%", "e": "0.000E+00", "2": "0.00", "k": "0.00"}.get(kind, "General")
            if "k" in kind:
                b.fill = kf
        row += 1
    for wtxt in r.warnings:
        ws.cell(row, 1, "⚠ " + wtxt).font = Font(name="Arial", color="C00000")
        row += 1

    d = wb.create_sheet("Data")
    out = raw.copy()
    for c in calc.columns:
        if c != "time":
            out[c] = calc[c].to_numpy()
    d.append(list(out.columns))
    for v in out.itertuples(index=False):
        d.append(list(v))
    for c in d[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = hf
    d.freeze_panes = "B2"
    n = len(out) + 1
    x = Reference(d, min_col=1, min_row=2, max_row=n)
    ch = ScatterChart()
    ch.title, ch.height, ch.width = "Fraction of N2O feed", 9, 18
    ch.x_axis.title, ch.y_axis.title = "Time (s)", "Fraction"
    ch.x_axis.delete = ch.y_axis.delete = False
    for name, color in (("f_n2o", "C00000"), ("f_n2", "2E75B6"), ("mass_balance", "7F7F7F")):
        if name in out.columns:
            ci = list(out.columns).index(name) + 1
            sr = Series(Reference(d, min_col=ci, min_row=1, max_row=n), x, title_from_data=True)
            sr.marker.symbol = "none"
            sr.graphicalProperties.line.solidFill = color
            sr.graphicalProperties.line.width = 12000
            ch.series.append(sr)
    ws.add_chart(ch, "E3")
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
